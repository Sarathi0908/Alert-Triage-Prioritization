"""
build_excels.py — the required 3-sheet Excel in data/ (minimum size):
  1_features_selected  — the native features given to the model, family, gain%, single-AUC
  2_features_removed   — every dropped column + WHY (near-constant noise / redundant noise / leak)
  3_label              — the label: attack_class -> y (0/1), counts, attack rate
"""
import sys, json
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.stdout.reconfigure(encoding="utf-8")

# This module lives at <slot>/src/reporting/ -> the slot root is three parents up.
_SLOT_ROOT = Path(__file__).resolve().parents[2]
if str(_SLOT_ROOT) not in sys.path:
    sys.path.insert(0, str(_SLOT_ROOT))
from src.config import DATA_DIR, SLOT_ROOT  # noqa: E402

ROOT = SLOT_ROOT
DATA = DATA_DIR
OUT = DATA / "CSE-CIC-IDS2018_feature_sheets.xlsx"

NAVY, BLUE, GREEN, GREENF, RED, REDF, AMBER, AMBERF = "1F3864","2E5496","C6EFCE","006100","FFC7CE","9C0006","FFEB9C","9C6500"
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(thin, thin, thin, thin)
WRAP = Alignment(wrap_text=True, vertical="top")

sel = json.load(open(DATA / "native_selected.json"))
scr = pd.read_csv(DATA / "feature_screen.csv")
man = json.load(open(DATA / "ids2018_feature_manifest.json"))

def hdr(ws, r, cols):
    for j, h in enumerate(cols, 1):
        c = ws.cell(r, j, h); c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=BLUE); c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center"); c.border = BORDER
    ws.row_dimensions[r].height = 26
def row(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = Font(size=9); c.alignment = WRAP; c.border = BORDER

wb = Workbook()

# 1 selected
s1 = wb.active; s1.title = "1_features_selected"
t = s1.cell(1, 1, f"MODEL INPUT — {sel['n']} NATIVE features (leak/noise-screened)"); t.font = Font(bold=True, color="FFFFFF", size=12)
t.fill = PatternFill("solid", fgColor=NAVY); s1.merge_cells("A1:D1")
hdr(s1, 2, ["feature", "family", "gain %", "single-feature AUC"])
for i, f in enumerate(sel["features"], 3):
    row(s1, i, [f, sel["families"][f], sel["gain_pct"][f], round(sel["single_feature_auc"][f], 3)])
    s1.cell(i, 1).fill = PatternFill("solid", fgColor=GREEN); s1.cell(i, 1).font = Font(size=9, bold=True, color=GREENF)
for col, w in {1: 30, 2: 16, 3: 10, 4: 18}.items(): s1.column_dimensions[get_column_letter(col)].width = w

# 2 removed
s2 = wb.create_sheet("2_features_removed")
t = s2.cell(1, 1, "REMOVED — why each column was dropped"); t.font = Font(bold=True, color="FFFFFF", size=12)
t.fill = PatternFill("solid", fgColor=NAVY); s2.merge_cells("A1:D1")
hdr(s2, 2, ["feature", "reason removed", "single-feature AUC", "detail"])
rm = sel["removed"]
reason_map = {}
for f in rm.get("near_constant_noise", []): reason_map[f] = ("NOISE — near-constant", "one value / >99.9% zero -> no information")
for f in rm.get("redundant_noise (|r|>0.95)", []): reason_map[f] = ("NOISE — redundant", "|r|>0.95 with a kept feature")
for f in rm.get("leak_suspect (auc>=0.999)", []): reason_map[f] = ("LEAK", "single-feature AUC >= 0.999")
for f in rm.get("not_in_2018_dataset (would be NOISE on 2018 test)", []): reason_map[f] = ("NOISE — not in 2018", "no equivalent column in the 2018 NetFlow dataset -> would be noise on the 2018 test")
for f in rm.get("dominant_removed_for_balance", []): reason_map[f] = ("DOMINANT (removed)", "held >22% of gain alone -> removed so importance is balanced (memorisation/leak risk)")
for f in rm.get("low_gain_tail", []): reason_map[f] = ("low-gain tail", "carried <2% gain -> dead weight, dropped for a balanced set")
r = 3
aucmap = dict(zip(scr.feature, scr.single_feature_auc))
for f, (why, detail) in sorted(reason_map.items()):
    row(s2, r, [f, why, round(float(aucmap.get(f, 0)), 3), detail])
    fill = RED if "LEAK" in why else AMBER; ff = REDF if "LEAK" in why else AMBERF
    s2.cell(r, 2).fill = PatternFill("solid", fgColor=fill); s2.cell(r, 2).font = Font(size=9, bold=True, color=ff)
    r += 1
for col, w in {1: 30, 2: 22, 3: 16, 4: 40}.items(): s2.column_dimensions[get_column_letter(col)].width = w

# 3 label
s3 = wb.create_sheet("3_label")
t = s3.cell(1, 1, "LABEL — attack_class -> y_is_attack (BENIGN=0, attack=1)"); t.font = Font(bold=True, color="FFFFFF", size=12)
t.fill = PatternFill("solid", fgColor=NAVY); s3.merge_cells("A1:C1")
hdr(s3, 3, ["attack_class", "y (0/1)", "count"])
ac = man["attack_classes"]
r = 4
for cls, n in sorted(ac.items(), key=lambda kv: -kv[1]):
    yv = 0 if cls.upper() == "BENIGN" else 1
    row(s3, r, [cls.replace("�", "-"), yv, n]); r += 1
s3.cell(2, 1, f"total rows {man['rows']:,} · attack rate {man['attack_rate_%']}%").font = Font(size=9, italic=True)
for col, w in {1: 30, 2: 10, 3: 14}.items(): s3.column_dimensions[get_column_letter(col)].width = w
for ws in wb.worksheets: ws.sheet_view.showGridLines = False

wb.save(OUT)
print(f"wrote {OUT}  (3 sheets: selected {sel['n']}, removed {len(reason_map)}, label {len(ac)} classes)")
